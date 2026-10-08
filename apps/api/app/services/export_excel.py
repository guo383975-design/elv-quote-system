"""报价单 → .xlsx 字节流。

只消费 `app.services.export_data.ExportContext`，不碰 ORM、不碰权限。

## 两个必须守住的细节

1. **金额写成数字，不写成文本。** ctx 里是 `"1,234.56"` 这种给人看的字符串，
   直接写进单元格的话甲方在 Excel 里没法求和（而且单元格带个绿三角）。
   这里把它解析回数字并设 `#,##0.00`，解析不了才原样当文本。

2. **空串必须是空单元格，绝不能写成 0。** ctx 用空串表示「这个计价口径里
   没有这一项」。写成 0 会让甲方理解成「这一项是零元」，那是两个完全不同的
   结论 —— 报价单上这种歧义是要出事的。
"""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.export_data import ExportContext

# ---------------------------------------------------------------------------
# 样式
# ---------------------------------------------------------------------------

_HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")
_TITLE_FONT = Font(bold=True, size=14)
_HEAD_FONT = Font(bold=True, size=10)
_LABEL_FONT = Font(bold=True, size=10)
_TOTAL_FONT = Font(bold=True, size=10)
_COST_FONT = Font(size=10, color="9C0006")  # 成本行用暗红，提醒这是内部数据

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right", vertical="center")
_CENTER = Alignment(horizontal="center", vertical="center")

_MONEY_FMT = "#,##0.00"
_QTY_FMT = "0.###"
_PCT_FMT = "0%"

_NUM_RE = re.compile(r"^-?[\d,]+(\.\d+)?$")


# ---------------------------------------------------------------------------
# 单元格值换算
# ---------------------------------------------------------------------------

def _num(text: str) -> float | str | None:
    """把 ctx 的展示字符串换成可参与计算的数字。

    返回 None 表示「空」→ 调用方写空单元格。
    返回 str 表示「不是数字」→ 原样写文本（比如「（本报价单暂无明细）」）。
    """
    if text is None:
        return None
    s = str(text).strip()
    if not s:
        return None
    if s.endswith("%"):
        body = s[:-1].strip()
        try:
            return float(Decimal(body) / Decimal(100))
        except (InvalidOperation, ArithmeticError):
            return s
    if _NUM_RE.match(s):
        return float(Decimal(s.replace(",", "")))
    return s


def _put(ws: Worksheet, row: int, col: int, *, text: str, fmt: str, align: Alignment,
         font: Font | None = None, fill: PatternFill | None = None, border: bool = True) -> None:
    """写一个「可能是数字」的单元格。"""
    c = ws.cell(row=row, column=col)
    val = _num(text)
    if val is None:
        # 关键：空串 → 真空单元格，不是 0
        c.value = None
    else:
        c.value = val
        if isinstance(val, (int, float)):
            c.number_format = fmt
    c.alignment = align
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if border:
        c.border = _BORDER


def _put_text(ws: Worksheet, row: int, col: int, text: str, *,
              font: Font | None = None, align: Alignment | None = None,
              fill: PatternFill | None = None, border: bool = True) -> None:
    c = ws.cell(row=row, column=col, value=text if text != "" else None)
    c.alignment = align or _LEFT
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if border:
        c.border = _BORDER


def _width(ws: Worksheet, col: int, text: str) -> None:
    """按内容估算列宽（中文算 2 个字符宽），夹在 8~46 之间。

    夹上界是必须的：备注里可能出现很长的一句话，不夹的话列宽会到几百，
    打开就是一片空白要往右拖半天。
    """
    if text is None:
        return
    s = str(text)
    width = sum(2 if ord(ch) > 0x2E7F else 1 for ch in s) + 2
    letter = get_column_letter(col)
    cur = ws.column_dimensions[letter].width or 0
    ws.column_dimensions[letter].width = max(cur, min(max(width, 8), 46))


def _print_setup(ws: Worksheet, *, last_col: int, title_rows: str, landscape: bool = True) -> None:
    """打印设置（对应需求里的「打印」）。"""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows          # 跨页重复表头
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------
# Sheet 1：报价汇总
# ---------------------------------------------------------------------------

def _sheet_summary(wb: Workbook, ctx: ExportContext) -> None:
    ws = wb.active
    ws.title = "报价汇总"
    m = ctx.meta

    ws.cell(row=1, column=1, value=m.title or "工程报价单").font = _TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    ws.cell(row=1, column=1).alignment = _LEFT
    ws.cell(row=2, column=1, value=ctx.meta.company_name or "").font = Font(size=10)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3)
    ws.cell(row=2, column=1).alignment = _LEFT

    info: list[tuple[str, str]] = [
        ("报价单号", m.quote_code),
    ]
    # 版本 = **业务版本的时间标记**（2026-10-08 用户要求「版本号以时间为标记」）✓；
    # 一次都没保存过 → version_label 为空 → **整行省略** ✓
    # （不能印「v0」/「v—」✗ —— 那是把"没保存过"说成了某个版本 ✗）
    if m.version_label:
        info.append(("版本", m.version_label))
    info += [
        ("状态", m.status_label),
        ("项目编号", m.project_code),
        ("项目名称", m.project_name),
        ("客户名称", m.customer_name),
        ("联系人", m.customer_contact),
        ("联系电话", m.customer_phone),
        ("项目地址", m.address),
        ("计税方式", m.tax_method_label),
        ("税率", m.tax_rate_pct),
        ("生成时间", m.generated_at),
    ]
    r = 4
    for label, value in info:
        _put_text(ws, r, 1, label, font=_LABEL_FONT, border=False)
        _put_text(ws, r, 2, value, border=False)
        r += 1

    r += 1
    head_row = r
    for i, h in enumerate(ctx.fee_headers(), start=1):
        _put_text(ws, r, i, h, font=_HEAD_FONT, align=_CENTER, fill=_HEAD_FILL)
    _width(ws, 1, "费用项")
    _width(ws, 2, "点数单价口径")
    _width(ws, 3, "定额计价口径")
    r += 1

    for fee in ctx.fees:
        cost_row = fee.label in ("成本合计", "毛利", "毛利率")
        font = _TOTAL_FONT if fee.emphasis else (_COST_FONT if cost_row else None)
        _put_text(ws, r, 1, fee.label, font=font or _LABEL_FONT)
        fmt = _PCT_FMT if fee.label == "毛利率" else _MONEY_FMT
        _put(ws, r, 2, text=fee.amount_point, fmt=fmt, align=_RIGHT, font=font)
        _put(ws, r, 3, text=fee.amount_quota, fmt=_MONEY_FMT, align=_RIGHT, font=font)
        # 金额列按内容撑宽（数字放不下会显示 ####），标签列也一并按内容撑宽
        _width(ws, 1, fee.label)
        _width(ws, 2, fee.amount_point)
        _width(ws, 3, fee.amount_quota)
        if fee.emphasis:
            for c in range(1, 4):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="DCE6F1")
        r += 1

    if m.notes:
        r += 1
        _put_text(ws, r, 1, "备注", font=_LABEL_FONT, border=False)
        _put_text(ws, r, 2, m.notes, border=False)

    ws.freeze_panes = f"A{head_row + 1}"
    _print_setup(ws, last_col=3, title_rows=f"{head_row}:{head_row}", landscape=False)


# ---------------------------------------------------------------------------
# Sheet 2：报价明细
# ---------------------------------------------------------------------------

def _sheet_items(wb: Workbook, ctx: ExportContext) -> None:
    ws = wb.create_sheet("报价明细")
    headers = ctx.item_headers()
    for i, h in enumerate(headers, start=1):
        _put_text(ws, 1, i, h, font=_HEAD_FONT, align=_CENTER, fill=_HEAD_FILL)
        _width(ws, i, h)

    if not ctx.items:
        _put_text(ws, 2, 1, "（本报价单暂无明细）", border=False)
    else:
        # 靠列名定位，避免硬编码列号（含成本时列会右移）
        idx = {name: i for i, name in enumerate(headers, start=1)}
        money_cols = {"单价", "合计", "进价", "成本小计", "毛利"}
        qty_cols = {"序号", "数量"}
        center_cols = {"单位", "来源"}
        for rowno, item in enumerate(ctx.items, start=2):
            for name, text in zip(headers, ctx.item_row_cells(item), strict=True):
                col = idx[name]
                if name in money_cols:
                    _put(ws, rowno, col, text=text, fmt=_MONEY_FMT, align=_RIGHT)
                    # 金额列也要按内容撑宽：数字放不下时 Excel 显示的是「####」，
                    # 不像文本那样溢出到邻列。只按表头（「单价」两个字）定宽的话，
                    # 遇到「1,256,269.56」这种大额就会被截成 ####。
                    _width(ws, col, text)
                elif name in qty_cols:
                    # 序号/数量写成数字：写成文本的话 Excel 里没法按序号排序，
                    # 而且单元格会带绿三角提示"以文本形式存储的数字"。
                    _put(ws, rowno, col, text=text, fmt=_QTY_FMT, align=_CENTER)
                elif name in center_cols:
                    _put_text(ws, rowno, col, text, align=_CENTER)
                else:
                    _put_text(ws, rowno, col, text)
                    _width(ws, col, text)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    _print_setup(ws, last_col=len(headers), title_rows="1:1")


# ---------------------------------------------------------------------------
# Sheet 3：子系统汇总
# ---------------------------------------------------------------------------

def _sheet_subsystems(wb: Workbook, ctx: ExportContext) -> None:
    ws = wb.create_sheet("子系统汇总")
    headers = ctx.subsystem_headers()
    for i, h in enumerate(headers, start=1):
        _put_text(ws, 1, i, h, font=_HEAD_FONT, align=_CENTER, fill=_HEAD_FILL)
        _width(ws, i, h)

    if not ctx.subsystems:
        _put_text(ws, 2, 1, "（本项目暂无子系统）", border=False)
    else:
        money_cols = {"设备费", "安装费", "辅材费", "小计"}
        for rowno, sub in enumerate(ctx.subsystems, start=2):
            for name, text in zip(headers, ctx.subsystem_row_cells(sub), strict=True):
                i = headers.index(name) + 1
                if name in money_cols:
                    _put(ws, rowno, i, text=text, fmt=_MONEY_FMT, align=_RIGHT)
                    _width(ws, i, text)   # 同明细表：金额列要按内容撑宽，否则大额显示 ####
                elif name in ("序号", "点位数"):
                    _put(ws, rowno, i, text=text, fmt=_QTY_FMT, align=_CENTER)
                else:
                    _put_text(ws, rowno, i, text)
                    _width(ws, i, text)

    ws.freeze_panes = "A2"
    _print_setup(ws, last_col=len(headers), title_rows="1:1")


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

def render(ctx: ExportContext) -> bytes:
    """把 `ExportContext` 渲染成 xlsx 字节。纯函数：不写盘、不打印。"""
    wb = Workbook()
    _sheet_summary(wb, ctx)
    _sheet_items(wb, ctx)
    _sheet_subsystems(wb, ctx)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["render"]

