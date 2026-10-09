"""每单「公司抬头」测试（2026-10-09）。

用户需求：有时要用**别家公司抬头**去报价 —— 只开放**公司名称**一个字段
（地址 / 电话 / 邮箱 / 税号 / Logo / 盖章位一律不做 ✗）。

本文件盯住两条硬约束：
  ① **老单不许变**：`company_name` 为空的老单，三处出口仍必须是「XX科技」——
     与加这个功能之前逐字一致；
  ② **快照性**（最容易漏的一条）：本单选完抬头后，**改掉抬头库**，那张单仍显示
     **原来那个名字**。因为存的"就是当时选定的那个字符串本身"，
     不是抬头库里的下标 / id ⇒ 一旦有人改成"拿 quotes.company_name
     去查抬头库再渲染"，改库就会**追溯改写已开出的报价单** ✗✗。

三个出口都**真读渲染产物**（不是只看"理论上走了收口"✗ —— 我们在「毛利泄露」
上吃过"以为走了其实没走"的亏）：
  · `export-clauses` 接口 —— 前端「客户版预览页」的**唯一**数据源 ✓
  · 客户版 PDF —— pypdf 把正文文字抽出来断言 ✓
  · Excel —— openpyxl 读「报价汇总」的 A2 单元格 ✓

（内部版 PDF 与客户版 PDF 是**同一个路由**的 `?internal=true` ✓，
抬头取自同一个 `meta.company_name` ✓，所以抽查客户版即可覆盖两者取值口径 ✓。）
"""

from __future__ import annotations

import io

from httpx import AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.project import Project, ProjectStatus, ProjectType

# 系统设置里 company.name 的出厂默认值（= 加这个功能之前所有单据印的那个）
DEFAULT_COMPANY = "XX科技有限公司"


# ---------------------------------------------------------------------------
# 助手
# ---------------------------------------------------------------------------

async def _project(db: AsyncSession, code: str = "PRJ-LH-1") -> Project:
    p = Project(
        code=code, name="抬头测试项目", customer_name="客户",
        project_type=ProjectType.COMMERCIAL, status=ProjectStatus.DRAFT,
    )
    db.add(p)
    await db.commit()
    await db.refresh(p)
    return p


async def _quote(
    client: AsyncClient, headers, project_id, code: str = "QT-LH-1"
) -> dict:
    r = await client.post(
        "/api/v1/quotes",
        headers=headers,
        json={
            "project_id": str(project_id), "code": code, "title": "抬头测试单",
            "fee_base": "device_aux_install", "auxiliary_base": "device",
            "auxiliary_rate": "0.06", "design_rate": "0.02", "warranty_rate": "0.03",
            "mgmt_profit_rate": "0.08", "tax_rate": "0.09",
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


async def _clauses_company(client: AsyncClient, headers, quote_id) -> str:
    """`export-clauses` 里的公司名 —— 客户预览页读的就是这个字段 ✓"""
    r = await client.get(f"/api/v1/quotes/{quote_id}/export-clauses", headers=headers)
    assert r.status_code == 200, r.text
    return r.json()["company"]["name"]


def _xlsx_company(content: bytes) -> str:
    """Excel「报价汇总」A2 = 抬头单元格（见 export_excel._sheet_summary）✓"""
    wb = load_workbook(io.BytesIO(content))
    return str(wb["报价汇总"].cell(row=2, column=1).value or "")


def _pdf_text(content: bytes) -> str:
    """客户版 PDF 的正文文字 ✓（中文走子集内嵌字体的 ToUnicode CMap ✓，实测可抽 ✓）"""
    reader = PdfReader(io.BytesIO(content))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


async def _export_all(client: AsyncClient, headers, quote_id) -> tuple[str, str, str]:
    """三处出口各取一次公司名：(预览接口, PDF正文, Excel A2) ✓"""
    preview = await _clauses_company(client, headers, quote_id)

    r = await client.get(f"/api/v1/quotes/{quote_id}/export.pdf", headers=headers)
    assert r.status_code == 200, r.text
    assert r.content[:5] == b"%PDF-", r.content[:5]
    pdf = _pdf_text(r.content)

    r = await client.get(f"/api/v1/quotes/{quote_id}/export.xlsx", headers=headers)
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK", r.content[:2]
    xlsx = _xlsx_company(r.content)

    return preview, pdf, xlsx


# ---------------------------------------------------------------------------
# ① 老单：没选抬头 → 三处仍是「XX科技」，且数据一个字节没被动过
# ---------------------------------------------------------------------------

async def test_legacy_quote_without_letterhead_still_prints_legacy_brand(
    client: AsyncClient, db: AsyncSession, admin_headers
) -> None:
    p = await _project(db)
    q = await _quote(client, admin_headers, p.id)

    # 新建的单没选抬头 = NULL ✓（不是被写成"XX科技"三个字 ✗ —— 那样以后改
    # 系统设置里的公司名，老单就不会跟着变了，等于改了老单行为 ✗）。
    assert q["company_name"] is None, q

    preview, pdf, xlsx = await _export_all(client, admin_headers, q["id"])
    assert preview == DEFAULT_COMPANY, preview
    assert DEFAULT_COMPANY in pdf, pdf[:300]
    assert xlsx == DEFAULT_COMPANY, xlsx


# ---------------------------------------------------------------------------
# ② 选了自定义抬头 → 三处都显示该名字
# ---------------------------------------------------------------------------

async def test_custom_letterhead_shows_in_all_three_outlets(
    client: AsyncClient, db: AsyncSession, admin_headers
) -> None:
    p = await _project(db)
    q = await _quote(client, admin_headers, p.id)
    custom = "浙江某某科技有限公司"

    r = await client.patch(
        f"/api/v1/quotes/{q['id']}", headers=admin_headers, json={"company_name": custom}
    )
    assert r.status_code == 200, r.text
    assert r.json()["company_name"] == custom, r.json()

    preview, pdf, xlsx = await _export_all(client, admin_headers, q["id"])
    assert preview == custom, preview
    assert custom in pdf, pdf[:300]
    assert xlsx == custom, xlsx
    # 系统设置里的公司名**没被改**（本单覆盖不污染全局）✓
    s = (await client.get("/api/v1/settings", headers=admin_headers)).json()
    assert s["company"]["name"] == DEFAULT_COMPANY, s["company"]


# ---------------------------------------------------------------------------
# ③ 快照性：本单选了之后**改掉抬头库**，那张单仍显示原来那个名字 ★核心
# ---------------------------------------------------------------------------

async def test_letterhead_is_snapshot_and_survives_library_change(
    client: AsyncClient, db: AsyncSession, admin_headers
) -> None:
    p = await _project(db)
    q = await _quote(client, admin_headers, p.id)
    a, b = "甲抬头有限公司", "乙抬头有限公司"

    # 抬头库里先有 A；本单选了 A
    r = await client.patch(
        "/api/v1/settings", headers=admin_headers, json={"company": {"letterheads": [a]}}
    )
    assert r.status_code == 200, r.text
    r = await client.patch(
        f"/api/v1/quotes/{q['id']}", headers=admin_headers, json={"company_name": a}
    )
    assert r.status_code == 200, r.text

    # ★ 之后把抬头库改掉：A 被删、只剩 B；**连系统设置里的公司名也改掉**
    #   （两个"以后可能被改"的地方一起动，才能证明本单真的锁住了 A）✗
    r = await client.patch(
        "/api/v1/settings",
        headers=admin_headers,
        json={"company": {"letterheads": [b], "name": "改过的总公司名"}},
    )
    assert r.status_code == 200, r.text
    assert (await client.get("/api/v1/settings", headers=admin_headers)).json()[
        "company"
    ]["name"] == "改过的总公司名"

    # 那张单**必须仍是 A** —— 三处都要 ✓
    preview, pdf, xlsx = await _export_all(client, admin_headers, q["id"])
    assert preview == a, f"到期接口被追溯改写了：{preview}"
    assert a in pdf, f"PDF 被追溯改写了：{pdf[:300]}"
    assert xlsx == a, f"Excel 被追溯改写了：{xlsx}"

    # 而且新单才用新的系统公司名 ✓ —— 证明"改设置"这条路本身是通的 ✓
    #（否则上面那条断言可能只是"设置根本没生效"造成的假绿 ✗）
    q2 = await _quote(client, admin_headers, p.id, code="QT-LH-2")
    assert await _clauses_company(client, admin_headers, q2["id"]) == "改过的总公司名"


# ---------------------------------------------------------------------------
# ④ 库里没有 / 被删掉的名字 + 空值：都不许崩
# ---------------------------------------------------------------------------

async def test_unknown_or_empty_letterhead_does_not_crash(
    client: AsyncClient, db: AsyncSession, admin_headers
) -> None:
    """两种边界都要**不崩**，且语义各不相同 ✓。

    ⚠ 口径（重要，别改坏）：本单**已经存了**一个非空名字时，即使抬头库里
      已经没有它，也要**照原样显示** ✓ —— 这正是"快照"的含义 ✓。
      若这里改成"回落默认"✗，就等于「删掉抬头库里的名字 → 已开出的单被改成XX科技」✗
      = **追溯改写历史单据** ✗✗，与硬约束 ② 直接冲突。
      「回落默认」只适用于**本单没存**（NULL / 空串）的情况 ✓。
    """
    p = await _project(db)
    q = await _quote(client, admin_headers, p.id)

    # (a) 库里根本没有这个名字 —— 不崩，且按存下来的原样显示 ✓
    ghost = "库里没有的某某工程有限公司"
    r = await client.patch(
        f"/api/v1/quotes/{q['id']}", headers=admin_headers, json={"company_name": ghost}
    )
    assert r.status_code == 200, r.text
    preview, pdf, xlsx = await _export_all(client, admin_headers, q["id"])
    assert preview == ghost, preview
    assert ghost in pdf, pdf[:300]
    assert xlsx == ghost, xlsx

    # (b) 空白串 → 视同"没选" → 回落系统设置 ✓（不崩，也不会印出一片空白 ✗）
    r = await client.patch(
        f"/api/v1/quotes/{q['id']}", headers=admin_headers, json={"company_name": "   "}
    )
    assert r.status_code == 200, r.text
    preview, pdf, xlsx = await _export_all(client, admin_headers, q["id"])
    assert preview == DEFAULT_COMPANY, preview
    assert DEFAULT_COMPANY in pdf, pdf[:300]
    assert xlsx == DEFAULT_COMPANY, xlsx

    # (c) 显式传 null = 恢复用系统设置 ✓（前端下拉选回「默认（系统设置）」就是这条路 ✓）
    r = await client.patch(
        f"/api/v1/quotes/{q['id']}", headers=admin_headers, json={"company_name": None}
    )
    assert r.status_code == 200, r.text
    assert r.json()["company_name"] is None, r.json()
    assert await _clauses_company(client, admin_headers, q["id"]) == DEFAULT_COMPANY


# ---------------------------------------------------------------------------
# ⑤ 抬头库本身的校验（净化规则：去空、去重、限长、限量、类型错 422）
# ---------------------------------------------------------------------------

async def test_letterheads_library_is_sanitised(
    client: AsyncClient, db: AsyncSession, admin_headers
) -> None:
    """净化规则 + 「只传 name 时不许抹掉抬头库」✓。

    ⚠ 结构在第二批（2026-10-09 用户补充「地址和联系电话和邮箱」）扩成了
      `[{name, address, phone, email}]` ✓ —— 旧格式（字符串数组）仍兼容 ✓，
      详见 `test_quote_letterhead_fields.py::test_letterheads_store_four_fields_and_normalise_legacy_strings` ✓。
      这里保留的是**净化/校验/不误抹**这几条 ✓。
    """
    # 去空 + 去重（保留首次出现的顺序）✓ —— 归一成四字段对象 ✓
    r = await client.patch(
        "/api/v1/settings",
        headers=admin_headers,
        json={"company": {"letterheads": [" 甲公司 ", "", "乙公司", "甲公司", "   "]}},
    )
    assert r.status_code == 200, r.text
    assert r.json()["company"]["letterheads"] == [
        {"name": "甲公司", "address": "", "phone": "", "email": ""},
        {"name": "乙公司", "address": "", "phone": "", "email": ""},
    ], r.json()

    # 不是数组 → 422
    r = await client.patch(
        "/api/v1/settings", headers=admin_headers, json={"company": {"letterheads": "甲公司"}}
    )
    assert r.status_code == 422, r.text

    # 数组里混进数字 → 422（既不是字符串也不是对象 = 接口误用）
    r = await client.patch(
        "/api/v1/settings",
        headers=admin_headers,
        json={"company": {"letterheads": ["甲公司", 123]}},
    )
    assert r.status_code == 422, r.text

    # 单条（名称）超长 → 422
    r = await client.patch(
        "/api/v1/settings",
        headers=admin_headers,
        json={"company": {"letterheads": ["甲" * 256]}},
    )
    assert r.status_code == 422, r.text

    # 超过上限 → 422
    r = await client.patch(
        "/api/v1/settings",
        headers=admin_headers,
        json={"company": {"letterheads": [f"公司{i}" for i in range(21)]}},
    )
    assert r.status_code == 422, r.text

    # 只传 name 时**不许**把 letterheads 抹掉 ✓（PATCH 是"只覆盖传入的键"）
    r = await client.patch(
        "/api/v1/settings", headers=admin_headers, json={"company": {"name": "新总公司"}}
    )
    assert r.status_code == 200, r.text
    assert r.json()["company"]["letterheads"] == [
        {"name": "甲公司", "address": "", "phone": "", "email": ""},
        {"name": "乙公司", "address": "", "phone": "", "email": ""},
    ], r.json()
