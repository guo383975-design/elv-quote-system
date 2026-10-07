"""导入接口测试（2026-10-03 起：设备导入 preview 要回「字段映射」）。"""

from __future__ import annotations

import pytest
from httpx import AsyncClient

# ---------------------------------------------------------------------------
# 原型第 6 页要的「字段映射」：preview 接口回 Excel 列 → 系统字段 → 示例值
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_device_preview_returns_field_mapping(
    client: AsyncClient, admin_headers, tmp_path
) -> None:
    """原型第 6 页左侧那栏「字段映射 Excel 列 → 系统字段」需要 preview 回 mapping。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    # 故意用同义词表头（规格/型号），验证同义词库识别 + 示例值
    ws.append(["设备名称", "品牌", "规格型号", "单位", "参考单价"])
    ws.append(["网络高清枪型摄像机", "海康威视", "DS-2CD2T46WD-I5", "台", 680])
    p = tmp_path / "dev.xlsx"
    wb.save(p)

    with p.open("rb") as fh:
        r = await client.post(
            "/api/v1/imports/devices/preview",
            files={"file": ("dev.xlsx", fh.read(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"duplicate_policy": "skip", "max_rows": "50"},
            headers=admin_headers,
        )
    assert r.status_code == 200, r.text
    body = r.json()
    assert "mapping" in body, body
    by_field = {m["system_field"]: m for m in body["mapping"]}
    assert by_field["name"]["excel_header"] == "设备名称"
    assert by_field["model"]["excel_header"] == "规格型号", body["mapping"]
    assert by_field["name"]["sample"] == "网络高清枪型摄像机"
    assert body["header_row"] >= 1


@pytest.mark.asyncio
async def test_inquiry_preview_returns_unit_and_candidates(
    client: AsyncClient, admin_headers, tmp_path
) -> None:
    """原型第 7 页：左栏「原文」要单位，右栏「候选设备」要 candidates（含分数）。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["甲方设备名称", "单位", "数量", "品牌要求"])
    ws.append(["网络高清枪型摄像机", "台", 10, "海康威视"])
    p = tmp_path / "inquiry.xlsx"
    wb.save(p)

    with p.open("rb") as fh:
        r = await client.post(
            "/api/v1/imports/inquiry/preview",
            files={"file": ("甲方询价单.xlsx", fh.read(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"threshold": "75", "limit": "20"},
            headers=admin_headers,
        )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["results"], body
    row = body["results"][0]
    assert row["unit"] == "台", row
    assert row["row"] == 2  # 第 1 行是表头
    assert "candidates" in row, row
    assert isinstance(row["candidates"], list)
    for c in row["candidates"]:
        assert "score" in c and "device_id" in c

async def test_device_template_xlsx_is_reimportable(client: AsyncClient, admin_headers) -> None:
    """M-19：下载的模板必须是**真 xlsx**，且填完能直接导回。

    以前这个按钮给的是前端拼的 .csv，而上传口只收 .xlsx —— 用户照模板填完传不回去。
    """
    from io import BytesIO

    from openpyxl import load_workbook

    r = await client.get("/api/v1/imports/device-template.xlsx", headers=admin_headers)
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK", "不是 xlsx"
    assert "spreadsheetml" in r.headers["content-type"]
    assert "filename*=UTF-8''" in r.headers["content-disposition"]

    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header[0] == "设备名称" and "规格型号" in header, header

    # 关键：把模板原样喂回导入预览，必须被识别（列名与解析器一致）
    r2 = await client.post(
        "/api/v1/imports/devices/preview",
        files={"file": ("设备库导入模板.xlsx", r.content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"duplicate_policy": "skip", "max_rows": "50"},
        headers=admin_headers,
    )
    assert r2.status_code == 200, r2.text
    mapped = {m["system_field"] for m in r2.json()["mapping"]}
    assert {"name", "model", "unit", "reference_price"} <= mapped, r2.json()["mapping"]


async def test_inquiry_template_xlsx_is_reimportable(client: AsyncClient, admin_headers) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    r = await client.get("/api/v1/imports/inquiry-template.xlsx", headers=admin_headers)
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK"
    wb = load_workbook(BytesIO(r.content))
    header = [c.value for c in wb.active[1]]
    assert header[0] == "设备名称", header

    r2 = await client.post(
        "/api/v1/imports/inquiry/preview",
        files={"file": ("甲方清单导入模板.xlsx", r.content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"threshold": "75", "limit": "20"},
        headers=admin_headers,
    )
    assert r2.status_code == 200, r2.text
