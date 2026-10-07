"""审计发现的越权：成本可见 ≠ 利润可见（2026-10-05）。

这两条来自"多账号权限审计"的实测发现，不是理论推演：

  ① `GET /quotes/{id}/cost` 原来只判「成本」，于是只开成本的账号能读到 `gross_profit`
  ② 内部版导出（`internal=true`）原来跟着"成本"一起带出「毛利 / 毛利率」两行

根因：2026-10-05 把可见性拆成 `data_cost` / `data_profit` 两档之后，
这两个出口漏了第二档。断言留着，防止再漏。
"""

from __future__ import annotations

import io

from httpx import AsyncClient
from openpyxl import load_workbook


async def _quote_with_item(client: AsyncClient, headers, code: str) -> dict:
    dev = (await client.post("/api/v1/devices", headers=headers, json={
        "code": f"LEAK-{code}", "name": "越权测试设备", "model": f"LEAK-{code}",
        "category": "video", "unit": "台", "reference_price": "880",
        "cost_price": "500"})).json()
    prj = (await client.post("/api/v1/projects", headers=headers, json={
        "code": f"LEAK-P-{code}", "name": "越权测试项目", "customer_name": "客户",
        "create_default_subsystems": False})).json()
    qt = (await client.post("/api/v1/quotes", headers=headers, json={
        "project_id": prj["id"], "code": f"LEAK-Q-{code}", "title": "越权测试单",
        "fee_base": "device_aux_install", "auxiliary_base": "device",
        "auxiliary_rate": "0", "design_rate": "0", "warranty_rate": "0",
        "mgmt_profit_rate": "0.2", "tax_rate": "0.09"})).json()
    # ⚠ name 是必填：漏了会 422，报价单里没有明细，测试就"假通过"了
    r_item = await client.post(f"/api/v1/quotes/{qt['id']}/items", headers=headers, json={
        "device_id": dev["id"], "name": "越权测试设备", "unit": "台",
        "quantity": "10", "unit_price": "880"})
    assert r_item.status_code in (200, 201), f"建明细失败：{r_item.status_code} {r_item.text[:120]}"
    await client.post(f"/api/v1/quotes/{qt['id']}/compute", headers=headers)
    return qt


async def test_cost_only_cannot_read_profit_from_cost_view(
    client: AsyncClient, admin_headers, est_headers
) -> None:
    """① 只给「成本」时，`/cost` 的毛利字段必须是 null（不是 0）。"""
    qt = await _quote_with_item(client, admin_headers, "CV")

    await client.patch("/api/v1/settings", headers=admin_headers, json={
        "permissions": {"overrides": {"estimator": {"data_cost": ["read"]}}}})
    try:
        r = await client.get(f"/api/v1/quotes/{qt['id']}/cost", headers=est_headers)
        assert r.status_code == 200, r.text
        body = r.json()
        # 成本照给
        assert body["cost_total"] is not None, body
        # 利润必须为 None —— 这是审计发现的越权点
        assert body["gross_profit"] is None, f"只开成本却拿到了毛利：{body}"
        assert body["gross_profit_rate"] is None, body
        assert body.get("profit_hidden_reason"), body

        # 再补上利润权限 → 这次应该有值
        await client.patch("/api/v1/settings", headers=admin_headers, json={
            "permissions": {"overrides": {"estimator": {
                "data_cost": ["read"], "data_profit": ["read"]}}}})
        r2 = await client.get(f"/api/v1/quotes/{qt['id']}/cost", headers=est_headers)
        assert r2.status_code == 200, r2.text
        assert r2.json()["gross_profit"] is not None, r2.json()
    finally:
        await client.patch("/api/v1/settings", headers=admin_headers,
                           json={"permissions": {"overrides": {}}})


async def test_cost_only_export_has_no_profit_rows(
    client: AsyncClient, admin_headers, est_headers
) -> None:
    """② 只给「成本」时，内部版导出里不该出现「毛利」行。"""
    qt = await _quote_with_item(client, admin_headers, "EX")

    await client.patch("/api/v1/settings", headers=admin_headers, json={
        "permissions": {"overrides": {"estimator": {"data_cost": ["read"]}}}})
    try:
        r = await client.get(f"/api/v1/quotes/{qt['id']}/export.xlsx",
                             headers=est_headers, params={"internal": "true"})
        assert r.status_code == 200, r.text
        wb = load_workbook(io.BytesIO(r.content))
        text = " ".join(
            str(cell.value)
            for ws in wb.worksheets for row in ws.iter_rows() for cell in row
            if cell.value is not None
        )
        assert "成本" in text, "有成本权限却没导出成本列"
        # 真有明细行，避免"报价单是空的"导致假通过
        assert "越权测试设备" in text, "导出里没有明细行（测试会假通过）"
        assert "毛利" not in text, "只开成本却导出了毛利行（审计发现的越权点）"

        # 补上利润权限 → 这次应该有毛利
        await client.patch("/api/v1/settings", headers=admin_headers, json={
            "permissions": {"overrides": {"estimator": {
                "data_cost": ["read"], "data_profit": ["read"]}}}})
        r2 = await client.get(f"/api/v1/quotes/{qt['id']}/export.xlsx",
                              headers=est_headers, params={"internal": "true"})
        wb2 = load_workbook(io.BytesIO(r2.content))
        text2 = " ".join(
            str(cell.value)
            for ws in wb2.worksheets for row in ws.iter_rows() for cell in row
            if cell.value is not None
        )
        assert "毛利" in text2, "补了利润权限却没导出毛利"
    finally:
        await client.patch("/api/v1/settings", headers=admin_headers,
                           json={"permissions": {"overrides": {}}})
