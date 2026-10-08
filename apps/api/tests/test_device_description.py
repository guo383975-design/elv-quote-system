"""设备描述 + 报价单"设备参数"栏（2026-10-05）。

用户需求（修正后）：「不需要单独设置参数接口，只需要有一个添加描述的地方就好。
在选择产品的时候鼠标放上去会显示参数或者直接显示在报价单的设备参数一栏。」

所以这里钉住四件事：
  ① 设备能存/改**描述**（多行，性能参数就写在这儿）；
  ② 加设备进报价单时，描述**自动带成**「设备参数」（快照）；
  ③ 显式传了设备参数就不覆盖（手工改的内容说了算）；
  ④ 导出件里「设备参数」这一栏**只在内部版**出现（客户版一律没有，与客户版 PDF 对齐）。
"""

from __future__ import annotations

import io

from httpx import AsyncClient
from openpyxl import load_workbook

PARAMS = "吞吐量 10 Gbps；并发连接 200 万；接口 8×GE + 4×SFP+；IPsec 隧道 2000 条"


async def _device(client: AsyncClient, headers, *, code: str, desc: str | None) -> dict:
    body = {"code": code, "name": f"设备{code}", "model": code, "category": "network",
            "unit": "台", "reference_price": "8000", "cost_price": "5000"}
    if desc is not None:
        body["description"] = desc
    r = await client.post("/api/v1/devices", headers=headers, json=body)
    assert r.status_code == 201, r.text
    return r.json()


async def test_device_description_round_trip(
    client: AsyncClient, admin_headers
) -> None:
    """① 描述能存能改，详情里带得回来。"""
    dev = await _device(client, admin_headers, code="DSC1", desc=PARAMS)
    assert dev["description"] == PARAMS, dev

    got = (await client.get(f"/api/v1/devices/{dev['id']}", headers=admin_headers)).json()
    assert got["description"] == PARAMS, got

    r = await client.patch(f"/api/v1/devices/{dev['id']}", headers=admin_headers,
                           json={"description": "改后的参数"})
    assert r.status_code == 200, r.text
    assert r.json()["description"] == "改后的参数", r.json()


async def test_item_spec_text_is_filled_from_device(
    client: AsyncClient, admin_headers
) -> None:
    """② 加设备进报价单 → 设备参数自动带上；③ 显式传了就不覆盖。"""
    dev = await _device(client, admin_headers, code="DSC2", desc=PARAMS)
    prj = (await client.post("/api/v1/projects", headers=admin_headers, json={
        "code": "DSC-P", "name": "描述测试项目", "customer_name": "客户",
        "create_default_subsystems": False})).json()
    qt = (await client.post("/api/v1/quotes", headers=admin_headers, json={
        "project_id": prj["id"], "code": "DSC-Q", "title": "描述测试单",
        "fee_base": "device_aux_install", "auxiliary_base": "device",
        "auxiliary_rate": "0", "design_rate": "0", "warranty_rate": "0",
        "mgmt_profit_rate": "0.15", "tax_rate": "0.09"})).json()

    r = await client.post(f"/api/v1/quotes/{qt['id']}/items", headers=admin_headers,
                          json={"device_id": dev["id"], "name": dev["name"],
                                "quantity": "2", "unit_price": "8000"})
    assert r.status_code == 201, r.text
    assert r.json()["spec_text"] == PARAMS, r.json()

    # 显式给了设备参数 → 保留调用方的（可能是针对本单的定制说明）
    r2 = await client.post(f"/api/v1/quotes/{qt['id']}/items", headers=admin_headers,
                           json={"device_id": dev["id"], "name": dev["name"],
                                 "quantity": "1", "unit_price": "8000",
                                 "spec_text": "本单特殊配置：双电源"})
    assert r2.status_code == 201, r2.text
    assert r2.json()["spec_text"] == "本单特殊配置：双电源", r2.json()

    # 手工录入（无设备）也能填设备参数
    r3 = await client.post(f"/api/v1/quotes/{qt['id']}/items", headers=admin_headers,
                           json={"name": "手动行", "quantity": "1", "unit_price": "100",
                                 "spec_text": "手动写的参数"})
    assert r3.status_code == 201, r3.text
    assert r3.json()["spec_text"] == "手动写的参数", r3.json()


async def test_export_spec_column_only_internal(
    client: AsyncClient, admin_headers
) -> None:
    """④「设备参数」栏只在**内部版** Excel 里出现，客户版一律没有。

    2026-10-07 改（安全修复 + 与 PDF 对齐）：客户版 Excel 以前带着「设备参数」列 ✗
    —— 客户版 PDF 早就去掉了 ✓。这条测试原先是反着写的（断言客户版**有**该列 ✗），
    功能改对了、测试没跟着改 ✗ → 现在改成同时钉住两侧 ✓：
      · 客户版（默认不带 internal）**没有**该列，也没有它的值 ✓
      · 内部版（?internal=true）**有**该列，且写的是那份快照 ✓
    """
    dev = await _device(client, admin_headers, code="DSC3", desc="吞吐 5 Gbps")
    prj = (await client.post("/api/v1/projects", headers=admin_headers, json={
        "code": "DSC-P3", "name": "导出测试项目", "customer_name": "客户",
        "create_default_subsystems": False})).json()
    qt = (await client.post("/api/v1/quotes", headers=admin_headers, json={
        "project_id": prj["id"], "code": "DSC-Q3", "title": "导出测试单",
        "fee_base": "device_aux_install", "auxiliary_base": "device",
        "auxiliary_rate": "0", "design_rate": "0", "warranty_rate": "0",
        "mgmt_profit_rate": "0.15", "tax_rate": "0.09"})).json()
    await client.post(f"/api/v1/quotes/{qt['id']}/items", headers=admin_headers,
                      json={"device_id": dev["id"], "name": dev["name"],
                            "quantity": "1", "unit_price": "8000"})
    await client.post(f"/api/v1/quotes/{qt['id']}/compute", headers=admin_headers)

    async def _sheet_text(*, internal: bool) -> str:
        url = f"/api/v1/quotes/{qt['id']}/export.xlsx"
        if internal:
            url += "?internal=true"
        r = await client.get(url, headers=admin_headers)
        assert r.status_code == 200, r.text
        wb = load_workbook(io.BytesIO(r.content))
        return " ".join(
            str(c.value)
            for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if c.value is not None
        )

    # 客户版：不带「设备参数」列，也不带它的值 ✓（这份是要发给甲方的 ✗）
    customer = await _sheet_text(internal=False)
    assert "设备参数" not in customer, "客户版 Excel 不该有「设备参数」栏"
    assert "吞吐 5 Gbps" not in customer, "客户版 Excel 不该带设备参数的值"

    # 内部版：照旧有该列，且写的是那份快照 ✓
    internal = await _sheet_text(internal=True)
    assert "设备参数" in internal, "内部版 Excel 里没有「设备参数」栏"
    assert "吞吐 5 Gbps" in internal, "内部版 Excel 里没写设备参数的值"


async def test_spec_column_only_when_present(
    client: AsyncClient, admin_headers
) -> None:
    """没填设备参数的单子，明细表**不该**多出这一列（否则所有单子都变宽）。

    实测过：无条件加列会把客户版 PDF 从 1 页挤到 2 页 —— 多数单子没填参数，
    不该为这个功能付排版代价。表头与单元格必须同条件，否则列数错位。
    """
    from app.services.export_data import ExportContext, ExportItemRow, ExportMeta

    meta = ExportMeta(
        quote_code="Q", title="t", version_label="2026-10-08 15:33", status_label="草稿",
        project_code="P",
        project_name="p", customer_name="c", customer_contact="", customer_phone="",
        address="", tax_method_label="", tax_rate_pct="9%", generated_at="now",
        include_cost=False,
    )
    plain = ExportContext(meta=meta, items=[
        ExportItemRow(seq=1, subsystem="", name="设备", brand="", model="M",
                      unit="台", quantity="1", unit_price="1", amount="1",
                      source_label="", note=""),
    ])
    assert "设备参数" not in plain.item_headers(), plain.item_headers()
    assert len(plain.item_headers()) == len(plain.item_row_cells(plain.items[0])), \
        "表头与单元格列数必须一致"

    # 客户版（include_cost=False）**永远不输出参数列** ——
    # 2026-10-05 用户要求 C ✓：横向 14 列已很紧 ✓，参数列会把名称/品牌/型号/价格挤成竖排 ✗。
    # 客户版 PDF：由后端传 show_spec_column=False ✓（2026-10-05 ✓）
    from dataclasses import replace as _replace
    customer_with_spec = ExportContext(meta=_replace(meta, show_spec_column=False), items=[
        ExportItemRow(seq=1, subsystem="", name="设备", brand="", model="M",
                      spec_text="吞吐 10Gbps", unit="台", quantity="1",
                      unit_price="1", amount="1", source_label="", note=""),
    ])
    assert "设备参数" not in customer_with_spec.item_headers(), (
        f"客户版不该出现设备参数列：{customer_with_spec.item_headers()}"
    )

    # 参数列只在**内部版**（含成本）里、且确实填了参数时才出现 ✓
    withspec = ExportContext(
        meta=_replace(meta, include_cost=True, show_spec_column=True), items=[
        ExportItemRow(seq=1, subsystem="", name="设备", brand="", model="M",
                      spec_text="吞吐 10Gbps", unit="台", quantity="1",
                      unit_price="1", amount="1", source_label="", note=""),
    ])
    assert "设备参数" in withspec.item_headers(), withspec.item_headers()
    assert "吞吐 10Gbps" in withspec.item_row_cells(withspec.items[0])
    assert len(withspec.item_headers()) == len(
        withspec.item_row_cells(withspec.items[0])
    ), "有参数时表头与单元格列数也必须一致"
